"""DOCX adapter (optional dependency: python-docx)."""

from __future__ import annotations

from pathlib import Path

from rag.adapters.base import AdapterError, ExtractedDocument


class DocxAdapter:
    name = "docx"

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() == ".docx"

    def extract(self, path: Path, *, document_id: str) -> ExtractedDocument:
        try:
            import docx
        except ImportError as exc:
            raise AdapterError(
                "docx adapter requires optional dependency python-docx (pip install '.[formats]')"
            ) from exc
        try:
            document = docx.Document(str(path))
            content = "\n".join(p.text for p in document.paragraphs if p.text).strip()
        except Exception as exc:
            raise AdapterError(f"docx adapter failed for {path.name}: {exc}") from exc
        return ExtractedDocument(
            document_id=document_id,
            path=document_id,
            content=content,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            language=None,
            metadata={"format": "docx"},
            adapter=self.name,
        )
